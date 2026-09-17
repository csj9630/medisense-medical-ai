"""DOCX/PPTX의 OOXML 구조를 직접 읽어 RAG용 텍스트와 이미지를 추출합니다."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
import re
from dataclasses import replace
from typing import Callable, Protocol

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RELATIONSHIP_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..contracts import (
    ExtractedDocument,
    OcrEngine,
    OcrLine,
    OcrProcessingConfig,
    ProgressCallback,
    ValidatedDocument,
)
from ..errors import DocumentProcessingError, OfficeExtractionError
from ..preprocessing import preprocess_image


@dataclass(frozen=True)
class OfficeImage:
    """Office 패키지에서 찾은 OCR 후보 이미지입니다."""

    content: bytes
    content_type: str
    label: str


@dataclass(frozen=True)
class OfficeContentUnit:
    """DOCX 문서 또는 PPTX 슬라이드 단위의 직접 추출 결과입니다."""

    title: str
    text: str
    images: list[OfficeImage] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedOfficeDocument:
    """중심 처리 서비스가 형식과 무관하게 소비하는 Office 파싱 결과입니다."""

    units: list[OfficeContentUnit]
    page_count: int | None
    document_type: str
    warnings: list[str] = field(default_factory=list)


class OfficeDocumentParser(Protocol):
    """중심 서비스가 구체적인 DOCX/PPTX 라이브러리를 알지 않게 하는 계약입니다."""

    def parse(self, document: ValidatedDocument) -> ParsedOfficeDocument: ...


class DirectOfficeDocumentParser:
    """python-docx와 python-pptx를 사용해 Office 문서를 직접 분석합니다."""

    def parse(self, document: ValidatedDocument) -> ParsedOfficeDocument:
        if document.file_type == "docx":
            return self._parse_docx(document.content)
        if document.file_type == "pptx":
            return self._parse_pptx(document.content)
        raise OfficeExtractionError("직접 분석을 지원하지 않는 Office 형식입니다.")

    def _parse_docx(self, content: bytes) -> ParsedOfficeDocument:
        try:
            source = Document(BytesIO(content))
            text_blocks: list[str] = []
            for block in source.iter_inner_content():
                if isinstance(block, Paragraph):
                    rendered = _render_docx_paragraph(block)
                elif isinstance(block, Table):
                    rendered = _render_table(block.rows)
                else:
                    rendered = ""
                if rendered:
                    text_blocks.append(rendered)

            warnings: list[str] = [
                "DOCX 구조를 직접 분석했습니다. 실제 페이지 수는 계산하지 않습니다."
            ]
            body = source.element.body

            text_box_texts = _unique_non_empty(
                _text_from_xml_node(node)
                for node in body.xpath(".//w:txbxContent")
            )
            if text_box_texts:
                text_blocks.append(
                    "### 텍스트 상자\n\n" + "\n\n".join(text_box_texts)
                )
                warnings.append("텍스트 상자는 문서 끝에 모아 추출해 화면 순서와 다를 수 있습니다.")

            inserted_texts = _unique_non_empty(
                _text_from_xml_node(node)
                for node in body.xpath(".//w:ins")
            )
            if inserted_texts:
                text_blocks.append(
                    "### 변경 추적 삽입 내용\n\n" + "\n\n".join(inserted_texts)
                )
                warnings.append("변경 추적 삽입 내용은 문서 끝에 별도로 추출했습니다.")
            if body.xpath(".//w:del"):
                warnings.append("변경 추적에서 삭제된 내용은 추출하지 않았습니다.")
            if body.xpath(".//w:altChunk"):
                warnings.append("외부 콘텐츠(altChunk)는 직접 추출에서 제외했습니다.")
            if body.xpath(".//w:object"):
                warnings.append("OLE 또는 임베디드 객체는 직접 추출에서 제외했습니다.")

            for heading, part_name in (
                ("각주", "/word/footnotes.xml"),
                ("미주", "/word/endnotes.xml"),
            ):
                story_text = _extract_docx_story_part(source, part_name)
                if story_text:
                    text_blocks.append(f"### {heading}\n\n{story_text}")

            images = _extract_docx_images(source)
            return ParsedOfficeDocument(
                units=[
                    OfficeContentUnit(
                        title="DOCX 문서",
                        text="\n\n".join(text_blocks).strip(),
                        images=images,
                    )
                ],
                page_count=None,
                document_type="docx_direct",
                warnings=_deduplicate(warnings),
            )
        except OfficeExtractionError:
            raise
        except (KeyError, TypeError, ValueError, etree.XMLSyntaxError) as exc:
            raise OfficeExtractionError(
                "DOCX 내부 구조에서 텍스트를 추출하지 못했습니다. "
                "Word에서 새 문서로 저장하거나 PDF로 업로드해 주세요."
            ) from exc
        except Exception as exc:
            raise OfficeExtractionError(
                "DOCX 문서를 직접 분석하지 못했습니다. "
                "Word에서 새 문서로 저장하거나 PDF로 업로드해 주세요."
            ) from exc

    def _parse_pptx(self, content: bytes) -> ParsedOfficeDocument:
        try:
            presentation = Presentation(BytesIO(content))
            units: list[OfficeContentUnit] = []
            warnings: list[str] = []

            for slide_number, slide in enumerate(presentation.slides, start=1):
                text_blocks: list[str] = []
                images: list[OfficeImage] = []
                slide_warnings: list[str] = []

                for shape in sorted(
                    slide.shapes,
                    key=lambda item: (
                        int(getattr(item, "top", 0) or 0),
                        int(getattr(item, "left", 0) or 0),
                    ),
                ):
                    shape_texts, shape_images, shape_warnings = _parse_pptx_shape(
                        shape,
                        slide_number,
                    )
                    text_blocks.extend(shape_texts)
                    images.extend(shape_images)
                    slide_warnings.extend(shape_warnings)

                if slide.has_notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame
                    notes_text = notes_frame.text.strip() if notes_frame else ""
                    if notes_text:
                        text_blocks.append(f"### 발표자 노트\n\n{notes_text}")

                units.append(
                    OfficeContentUnit(
                        title=f"슬라이드 {slide_number}",
                        text="\n\n".join(_deduplicate(text_blocks)).strip(),
                        images=_deduplicate_images(images),
                    )
                )
                warnings.extend(slide_warnings)

            if not units:
                warnings.append("PPTX에 분석할 슬라이드가 없습니다.")
            warnings.append(
                "PPTX 도형을 위에서 아래, 왼쪽에서 오른쪽 순서로 직접 분석했습니다."
            )
            return ParsedOfficeDocument(
                units=units,
                page_count=len(units),
                document_type="pptx_direct",
                warnings=_deduplicate(warnings),
            )
        except OfficeExtractionError:
            raise
        except (KeyError, TypeError, ValueError, etree.XMLSyntaxError) as exc:
            raise OfficeExtractionError(
                "PPTX 내부 구조에서 텍스트를 추출하지 못했습니다. "
                "PowerPoint에서 새 문서로 저장하거나 PDF로 업로드해 주세요."
            ) from exc
        except Exception as exc:
            raise OfficeExtractionError(
                "PPTX 문서를 직접 분석하지 못했습니다. "
                "PowerPoint에서 새 문서로 저장하거나 PDF로 업로드해 주세요."
            ) from exc


def process_office_document(
    document: ValidatedDocument,
    config: OcrProcessingConfig,
    ocr_service_factory: Callable[[], OcrEngine],
    progress_callback: ProgressCallback | None = None,
    office_parser: OfficeDocumentParser | None = None,
) -> ExtractedDocument:
    """Office 직접 Text와 포함 Image OCR을 문서 단위 순서대로 합칩니다."""

    parser = office_parser or DirectOfficeDocumentParser()
    _report_progress(
        progress_callback,
        "parsing_office",
        22,
        f"{document.file_type.upper()} 문서 구조를 직접 분석하고 있습니다.",
    )
    parsed = parser.parse(document)
    _report_progress(
        progress_callback,
        "parsing_office",
        35,
        "Office 문서의 텍스트와 이미지 구조를 확인했습니다.",
    )

    unit_texts: list[str] = []
    raw_unit_texts: list[str] = []
    warnings = list(parsed.warnings)
    confidences: list[float] = []
    lines: list[OcrLine] = []
    ocr_image_count = 0
    total_images = sum(len(unit.images) for unit in parsed.units)
    processed_images = 0

    for unit_index, unit in enumerate(parsed.units):
        unit_parts = [f"## {unit.title}"]
        raw_unit_parts = [f"## {unit.title}"]
        if unit.text:
            unit_parts.append(unit.text)
            raw_unit_parts.append(unit.text)

        for image_index, image in enumerate(unit.images, start=1):
            progress = 35
            if total_images:
                progress += round(processed_images / total_images * 45)
            processed_images += 1
            try:
                processed = preprocess_image(
                    image.content,
                    config.max_image_side,
                    config.max_image_pixels,
                    enable_denoise=config.enable_denoise,
                    enable_deskew=config.enable_deskew,
                )
                if processed.width * processed.height < 4_096:
                    warnings.append(f"{image.label}가 너무 작아 OCR에서 제외했습니다.")
                    continue

                _report_progress(
                    progress_callback,
                    "loading_model",
                    progress,
                    f"{image.label}의 텍스트를 OCR하고 있습니다.",
                )
                ocr_result = ocr_service_factory().extract_text(processed)
                ocr_image_count += 1
                if ocr_result.line_count:
                    confidences.append(ocr_result.confidence)
                lines.extend(
                    replace(line, page=unit_index)
                    for line in ocr_result.lines
                )
                raw_ocr_text = ocr_result.raw_text or ocr_result.text
                if raw_ocr_text:
                    raw_unit_parts.append(
                        f"### 이미지 OCR {image_index}\n\n{raw_ocr_text.strip()}"
                    )
                if ocr_result.text:
                    unit_parts.append(
                        f"### 이미지 OCR {image_index}\n\n{ocr_result.text.strip()}"
                    )
                else:
                    warnings.append(
                        f"{image.label}에서 Confidence 기준을 통과한 텍스트를 찾지 못했습니다."
                    )
            except DocumentProcessingError as exc:
                warnings.append(f"{image.label} OCR을 건너뛰었습니다: {exc}")

        if len(unit_parts) > 1:
            unit_texts.append("\n\n".join(unit_parts))
        if len(raw_unit_parts) > 1:
            raw_unit_texts.append("\n\n".join(raw_unit_parts))

    extracted_text = "\n\n".join(unit_texts)
    raw_text = "\n\n".join(raw_unit_texts)
    average_confidence = (
        sum(confidences) / len(confidences)
        if confidences
        else (1.0 if extracted_text else 0.0)
    )
    _report_progress(
        progress_callback,
        "extracting",
        80,
        "Office 문서의 직접 추출을 완료했습니다.",
    )
    return ExtractedDocument(
        text=extracted_text,
        page_count=parsed.page_count,
        document_type=parsed.document_type,
        ocr_image_count=ocr_image_count,
        average_confidence=average_confidence,
        warnings=warnings,
        lines=lines,
        raw_text=raw_text,
    )


def _render_docx_paragraph(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""

    style_name = paragraph.style.name if paragraph.style is not None else ""
    heading_match = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
    if heading_match:
        level = max(1, min(int(heading_match.group(1)), 6))
        return f"{'#' * level} {text}"
    if "List Bullet" in style_name:
        return f"- {text}"
    if "List Number" in style_name:
        return f"1. {text}"
    return text


def _render_table(rows) -> str:
    values = [
        [_escape_table_cell(cell.text) for cell in row.cells]
        for row in rows
    ]
    if not values:
        return ""

    column_count = max(len(row) for row in values)
    normalized = [row + [""] * (column_count - len(row)) for row in values]
    header = normalized[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(column_count)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _escape_table_cell(value: str) -> str:
    return " ".join(value.split()).replace("|", "\\|")


def _extract_docx_images(source) -> list[OfficeImage]:
    images: list[OfficeImage] = []
    for relationship in source.part.rels.values():
        if relationship.reltype != DOCX_RELATIONSHIP_TYPE.IMAGE:
            continue
        target = relationship.target_part
        images.append(
            OfficeImage(
                content=target.blob,
                content_type=target.content_type,
                label="DOCX 이미지",
            )
        )
    return _deduplicate_images(images)


def _extract_docx_story_part(source, part_name: str) -> str:
    for part in source.part.package.parts:
        if str(part.partname) != part_name:
            continue
        root = etree.fromstring(part.blob)
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        story_blocks: list[str] = []
        element_name = "footnote" if "footnotes" in part_name else "endnote"
        for element in root.xpath(f"//w:{element_name}", namespaces=namespaces):
            raw_id = element.get(f"{{{namespaces['w']}}}id", "-1")
            try:
                if int(raw_id) <= 0:
                    continue
            except ValueError:
                continue
            text = _text_from_xml_node(element, namespaces)
            if text:
                story_blocks.append(text)
        return "\n\n".join(story_blocks)
    return ""


def _parse_pptx_shape(shape, slide_number: int):
    texts: list[str] = []
    images: list[OfficeImage] = []
    warnings: list[str] = []
    shape_type = getattr(shape, "shape_type", None)

    if shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in sorted(
            shape.shapes,
            key=lambda item: (
                int(getattr(item, "top", 0) or 0),
                int(getattr(item, "left", 0) or 0),
            ),
        ):
            child_texts, child_images, child_warnings = _parse_pptx_shape(
                child,
                slide_number,
            )
            texts.extend(child_texts)
            images.extend(child_images)
            warnings.extend(child_warnings)
        return texts, images, warnings

    if getattr(shape, "has_text_frame", False):
        text = _pptx_text_frame_text(shape.text_frame)
        if text:
            texts.append(text)

    if getattr(shape, "has_table", False):
        texts.append(_render_table(shape.table.rows))

    if shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE} or hasattr(
        shape,
        "image",
    ):
        try:
            image = shape.image
            images.append(
                OfficeImage(
                    content=image.blob,
                    content_type=image.content_type,
                    label=f"슬라이드 {slide_number} 이미지",
                )
            )
        except (AttributeError, KeyError, ValueError):
            warnings.append(f"슬라이드 {slide_number}의 연결 이미지를 읽지 못했습니다.")

    if getattr(shape, "has_chart", False):
        chart_text = _pptx_chart_text(shape.chart)
        if chart_text:
            texts.append(chart_text)
        warnings.append(f"슬라이드 {slide_number}의 차트는 제목과 계열명만 추출했습니다.")

    shape_type_name = getattr(shape_type, "name", str(shape_type))
    if shape_type_name in {"IGX_GRAPHIC", "DIAGRAM", "EMBEDDED_OLE_OBJECT", "LINKED_OLE_OBJECT"}:
        warnings.append(
            f"슬라이드 {slide_number}의 {shape_type_name} 요소는 완전하게 추출되지 않을 수 있습니다."
        )
    return texts, images, warnings


def _pptx_text_frame_text(text_frame) -> str:
    paragraphs: list[str] = []
    for paragraph in text_frame.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.level:
            text = f"{'  ' * paragraph.level}- {text}"
        paragraphs.append(text)
    return "\n".join(paragraphs)


def _pptx_chart_text(chart) -> str:
    values: list[str] = []
    if chart.has_title:
        title = chart.chart_title.text_frame.text.strip()
        if title:
            values.append(title)
    series_names = [str(series.name).strip() for series in chart.series if str(series.name).strip()]
    if series_names:
        values.append("차트 계열: " + ", ".join(series_names))
    return "\n".join(values)


def _text_from_xml_node(node, namespaces: dict[str, str] | None = None) -> str:
    if namespaces is None:
        texts = node.xpath(".//w:t/text()")
    else:
        texts = node.xpath(".//w:t/text()", namespaces=namespaces)
    return " ".join(str(text).strip() for text in texts if str(text).strip())


def _unique_non_empty(values) -> list[str]:
    return _deduplicate(value.strip() for value in values if value and value.strip())


def _deduplicate(values) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def _deduplicate_images(images: list[OfficeImage]) -> list[OfficeImage]:
    result: list[OfficeImage] = []
    seen: set[bytes] = set()
    for image in images:
        if image.content in seen:
            continue
        seen.add(image.content)
        result.append(image)
    return result


def _report_progress(
    callback: ProgressCallback | None,
    stage: str,
    progress: int,
    message: str,
) -> None:
    if callback is not None:
        callback(stage, max(0, min(progress, 99)), message)
