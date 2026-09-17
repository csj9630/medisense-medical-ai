"""파일을 저장하지 않고 재현 가능한 OCR Test 문서를 생성합니다."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
from docx import Document as WordDocument
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def make_office_package(file_type: str, include_image: bool = False) -> bytes:
    output = BytesIO()
    if file_type == "docx":
        document = WordDocument()
        document.add_heading("DOCX direct extraction heading", level=1)
        document.add_paragraph("DOCX paragraph for RAG extraction")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Diagnosis"
        table.cell(1, 1).text = "Sample"
        if include_image:
            document.add_picture(BytesIO(make_png()), width=Inches(2))
        document.save(output)
    elif file_type == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        text_box.text = "PPTX direct extraction title"
        slide.notes_slide.notes_text_frame.text = "PPTX speaker notes"
        if include_image:
            slide.shapes.add_picture(
                BytesIO(make_png()),
                Inches(1),
                Inches(2),
                width=Inches(4),
            )
        presentation.save(output)
    else:
        raise ValueError(f"지원하지 않는 테스트 Office 형식: {file_type}")
    return output.getvalue()


def make_minimal_unreadable_office_package(file_type: str) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?><Types '
            'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
        )
        if file_type == "docx":
            archive.writestr(
                "word/document.xml",
                '<?xml version="1.0" encoding="UTF-8"?><w:document '
                'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body/></w:document>",
            )
        elif file_type == "pptx":
            archive.writestr(
                "ppt/presentation.xml",
                '<?xml version="1.0" encoding="UTF-8"?><p:presentation '
                'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
            )
        else:
            raise ValueError(f"지원하지 않는 테스트 Office 형식: {file_type}")
    return output.getvalue()


def make_png(width: int = 500, height: int = 300) -> bytes:
    output = BytesIO()
    Image.new("RGB", (width, height), "white").save(output, format="PNG")
    return output.getvalue()


def make_jpg(width: int = 500, height: int = 300) -> bytes:
    output = BytesIO()
    Image.new("RGB", (width, height), "white").save(output, format="JPEG")
    return output.getvalue()


def make_rotated_jpg() -> bytes:
    output = BytesIO()
    image = Image.new("RGB", (400, 200), "white")
    exif = Image.Exif()
    exif[274] = 6
    image.save(output, format="JPEG", exif=exif)
    return output.getvalue()


def make_digital_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Native digital PDF text with enough characters for direct extraction.",
    )
    content = document.tobytes()
    document.close()
    return content


def make_scanned_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=500, height=300)
    page.insert_image(page.rect, stream=make_png())
    content = document.tobytes()
    document.close()
    return content


def make_multi_page_digital_pdf() -> bytes:
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "First page native text with enough characters for direct extraction.",
    )
    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "Second page native text with enough characters for direct extraction.",
    )
    content = document.tobytes()
    document.close()
    return content


def make_encrypted_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Password protected PDF")
    content = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-password",
        user_pw="user-password",
    )
    document.close()
    return content


def make_hybrid_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_text(
        (50, 60),
        "Native text before the embedded image with enough characters to use hybrid mode.",
    )
    page.insert_image(pymupdf.Rect(50, 100, 550, 400), stream=make_png())
    page.insert_text(
        (50, 450),
        "Native text after the embedded image keeps the original document order.",
    )
    content = document.tobytes()
    document.close()
    return content
